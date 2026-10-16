# -*- coding: utf-8 -*-
"""fill_zkhp_contract.py — 根据结算单 docx + ZKHP 合同模板生成填好数据的《工矿机械产品买卖合同》。

用法：
    python fill_zkhp_contract.py 结算单.docx 模板.doc [选项]

选项：
    -o, --output-dir  输出目录（默认当前目录）
    --project         项目简称（如"绿洲15"），写入产品表备注；缺省时从结算单文件名提取
    --model           型号规格（如"YTS15-W-Z"）；缺省时由项目简称末尾数字推导（YTS{数字}-W-Z），
                      推不出则保留模板原型号
    --date            签订日期 YYYY-MM-DD（默认今天）；合同编号自动改为 ZKHP+YYYYMMDD
    --keep-number     保留模板原合同编号，不随签订日期更新

总价 = 加工表合计 + 单算表合计，按模板惯例四舍五入取整。
"""
import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import _Cell
from lxml import etree

SOFFICE_FALLBACK = r"C:\Program Files\LibreOffice\program\soffice.com"


# ---------------------------------------------------------------- 基础工具

def find_soffice():
    return shutil.which("soffice") or (SOFFICE_FALLBACK if Path(SOFFICE_FALLBACK).exists() else None)


def convert_doc_to_docx(doc_path: Path, outdir: Path) -> Path:
    """.doc -> .docx，返回转换后路径；输入已是 .docx 则直接返回。"""
    if doc_path.suffix.lower() == ".docx":
        return doc_path
    soffice = find_soffice()
    if not soffice:
        sys.exit("找不到 LibreOffice soffice，无法转换 .doc 模板")
    subprocess.run([soffice, "--headless", "--convert-to", "docx", str(doc_path), "--outdir", str(outdir)],
                   check=True, capture_output=True)
    out = outdir / (doc_path.stem + ".docx")
    if not out.exists():
        sys.exit(f"模板转换失败：{doc_path}")
    return out


def set_cell_text(cell: _Cell, text: str):
    """写入单元格文本，保留首段首 run 的格式（字体/颜色/高亮/下划线），\n 转为 <w:br/>。"""
    tc = cell._tc
    ps = tc.findall(qn("w:p"))
    first = ps[0]
    for extra in ps[1:]:
        tc.remove(extra)
    runs = first.findall(qn("w:r"))
    if runs:
        keep = runs[0]
        for r in runs[1:]:
            first.remove(r)
        for child in list(keep):
            if child.tag != qn("w:rPr"):
                keep.remove(child)
    else:
        keep = etree.SubElement(first, qn("w:r"))
    for i, ln in enumerate(text.split("\n")):
        if i > 0:
            etree.SubElement(keep, qn("w:br"))
        t = etree.SubElement(keep, qn("w:t"))
        t.set(qn("xml:space"), "preserve")
        t.text = ln


def tr_tcs(tr):
    return tr.findall(qn("w:tc"))


def row_texts(row) -> list[str]:
    return [_Cell(tc, None).text for tc in tr_tcs(row._tr)]


def fill_tr(tr, texts: list[str]):
    """按顺序把 texts 写入 tr 的各 tc（保留格式），texts 不足则多余的 tc 置空。"""
    tcs = tr_tcs(tr)
    for i, tc in enumerate(tcs):
        set_cell_text(_Cell(tc, None), texts[i] if i < len(texts) else "")


# ---------------------------------------------------------------- 中文大写金额（整数）

def rmb_upper(n: int) -> str:
    """整数人民币大写，如 508551 -> 伍拾万捌仟伍佰伍拾壹元整"""
    digits = "零壹贰叁肆伍陆柒捌玖"
    units = ["", "拾", "佰", "仟"]
    big = ["", "万", "亿", "万亿"]
    if n == 0:
        return "零元整"
    s = str(n)
    groups = []
    while s:
        groups.append(s[-4:])
        s = s[:-4]
    groups.reverse()
    parts = []
    for gi, g in enumerate(groups):
        big_unit = big[len(groups) - 1 - gi]
        seg = []
        zero_pending = False
        for ci, ch in enumerate(g):
            d = int(ch)
            pos = len(g) - 1 - ci
            if d == 0:
                zero_pending = True
            else:
                if zero_pending:
                    seg.append("零")
                    zero_pending = False
                seg.append(digits[d] + units[pos])
        text = "".join(seg)
        if text:
            parts.append(text + big_unit)
        elif big_unit == "万" and parts:
            parts.append("")  # 占位，下面统一处理零
    out = ""
    prev_empty = False
    for p in parts:
        if p == "":
            prev_empty = True
            continue
        if prev_empty and not out.endswith("零"):
            out += "零"
        out += p
        prev_empty = False
    return (out or "零") + "元整"


# ---------------------------------------------------------------- 结算单解析

def find_settlement_tables(doc):
    """在结算单中定位加工表（11 列）和单算表（含 名称|单价|数量|总价|备注 表头的 5 列表）。"""
    jiagong, dan_suan = None, None
    for tb in doc.tables:
        if len(tb.columns) >= 10 and jiagong is None:
            jiagong = tb
        elif len(tb.columns) == 5:
            for r in tb.rows[:3]:
                texts = [c.text.strip() for c in r.cells]
                if "名称" in texts and "单价" in texts and "总价" in texts:
                    dan_suan = tb
                    break
    if jiagong is None or dan_suan is None:
        sys.exit("结算单中未找到加工表（≥10列）或单算表（5列含 名称/单价/总价 表头）")
    return jiagong, dan_suan


def parse_total(text: str) -> Decimal:
    return Decimal(re.sub(r"[^\d.\-]", "", text))


def table_total(tb, price_col: int) -> Decimal:
    """取含"合计"行中指定列的数值。"""
    for row in reversed(tb.rows):
        cells = row.cells
        if any("合计" in c.text for c in cells):
            for c in (cells[price_col], *reversed(cells)):
                m = re.search(r"-?[\d,]+(?:\.\d+)?", c.text)
                if m:
                    return parse_total(m.group(0))
    sys.exit("表中未找到合计行")


def classify_jiagong_rows(tb):
    """加工表正文行分组：返回 (body_groups, total_row)。跳过表头行。

    行类型：sub=子表头(第3列为'重量')，data=数据行(第3列为数值)，blank=空延续行。
    分组规则：sub 开启新组；data/blank 归入序号相同的当前组，否则自成一组。
    """
    rows = tb.rows
    header, total_row = rows[0], None
    body = []
    for row in rows[1:]:
        texts = row_texts(row)
        if any("合计" in t for t in texts):
            total_row = row
            continue
        body.append((row, texts))
    groups = []
    current = None
    for row, texts in body:
        col2 = texts[2].strip() if len(texts) > 2 else ""
        kind = "sub" if col2 == "重量" else ("data" if re.search(r"\d", col2) else "blank")
        if kind == "sub" or current is None or \
           (current and texts[0].strip() and current[0][1][0].strip() != texts[0].strip()):
            current = []
            groups.append(current)
        current.append((kind, texts))
    return groups, (row_texts(total_row) if total_row is not None else None)


def dansuan_data_rows(tb):
    """单算表：返回 (header_texts, data_texts_list, total_texts)。自动跳过标题行。"""
    rows = tb.rows
    header_idx = None
    for i, row in enumerate(rows[:3]):
        texts = [c.text.strip() for c in row.cells]
        if "名称" in texts and "总价" in texts:
            header_idx = i
            break
    if header_idx is None:
        sys.exit("单算表缺少表头行")
    data, total_texts = [], None
    for row in rows[header_idx + 1:]:
        texts = row_texts(row)
        if any("合计" in t for t in texts):
            total_texts = texts
        elif any(t.strip() for t in texts):
            data.append(texts)
    return row_texts(rows[header_idx]), data, total_texts


# ---------------------------------------------------------------- 模板填充

def replace_in_paragraph(paragraph, pattern: str, repl: str) -> bool:
    """在段落各 run 内正则替换；跨 run 未命中时改写最后一个非空 run。"""
    for r in paragraph.runs:
        if re.search(pattern, r.text):
            r.text = re.sub(pattern, repl, r.text)
            return True
    for r in reversed(paragraph.runs):
        if r.text.strip():
            r.text = re.sub(r"[\dA-Za-z\-]+$", repl, r.text)
            return True
    return False


def fill_product_table(tb, model, total_int, project):
    """产品表：行1=产品行，行2=人民币合并行。"""
    old = {}
    prod_row = tb.rows[1]
    old["model"] = prod_row.cells[1].text.strip()
    old["price"] = prod_row.cells[4].text.strip()
    old["remark"] = prod_row.cells[6].text.strip()
    if model:
        set_cell_text(prod_row.cells[1], model)
    set_cell_text(prod_row.cells[4], str(total_int))
    set_cell_text(prod_row.cells[5], str(total_int))
    if project:
        set_cell_text(prod_row.cells[6], project)
    # 人民币行（含"小写"的合并行；横向合并单元格的多个 cell 代理指向同一 tc，写首个即可）
    for row in tb.rows[2:]:
        for c in row.cells:
            if "小写" in c.text:
                set_cell_text(c, f"（小写）￥{total_int}元；（大写）{rmb_upper(total_int)}；（含13%增值税）；")
                return old
    sys.exit("产品表中未找到人民币（小写/大写）行")


def rebuild_jiagong(tpl_tb, groups, total_texts):
    """用模板原型行重建加工表正文：对子表头/数据/空行/单行/合计分别克隆原型并填文本。"""
    tbl = tpl_tb._tbl
    trs = tpl_tb._tbl.findall(qn("w:tr"))
    proto_sub = copy.deepcopy(trs[1])    # 子表头行（重量|气囊组件|...）
    proto_data = copy.deepcopy(trs[2])   # 数据行
    # 空延续行：找一行第3列及之后全空的正文行，找不到就用数据行
    proto_blank = None
    proto_single = None
    proto_total = copy.deepcopy(trs[-1])
    for tr in trs[3:-1]:
        texts = [_Cell(tc, None).text for tc in tr_tcs(tr)]
        if len(texts) >= 11 and texts[2].strip() and texts[2].strip() != "重量" \
           and not any("合计" in t for t in texts) and proto_single is None \
           and not texts[3].strip() and not texts[4].strip():
            # 单行项目（如"砂石斗提机平台"：非称重计价各列为空）
            if len(tr_tcs(tr)) == 11 and not tr_tcs(tr)[0].find(qn("w:tcPr")).find(qn("w:vMerge")):
                proto_single = copy.deepcopy(tr)
        if proto_blank is None and len(texts) >= 11 and all(not t.strip() for t in texts[2:]):
            proto_blank = copy.deepcopy(tr)
    if proto_single is None:
        proto_single = copy.deepcopy(proto_data)
    if proto_blank is None:
        proto_blank = copy.deepcopy(proto_data)

    new_trs = []
    for group in groups:
        if group[0][0] == "sub":
            tr = copy.deepcopy(proto_sub)
            fill_tr(tr, group[0][1])
            new_trs.append(tr)
            for kind, texts in group[1:]:
                tr = copy.deepcopy(proto_data if kind == "data" else proto_blank)
                fill_tr(tr, texts)
                new_trs.append(tr)
        else:
            for kind, texts in group:
                tr = copy.deepcopy(proto_single if kind == "data" else proto_blank)
                fill_tr(tr, texts)
                new_trs.append(tr)
    tr = copy.deepcopy(proto_total)
    if total_texts:
        fill_tr(tr, total_texts)
    new_trs.append(tr)

    # 用新行替换表头之后、合计行（含）之前的所有旧行
    old_trs = tbl.findall(qn("w:tr"))
    anchor = old_trs[0]
    for tr in old_trs[1:]:
        tbl.remove(tr)
    for tr in new_trs:
        tbl.append(tr)


def rebuild_dansuan(tpl_tb, data_rows, total_texts):
    """单算表：按结算单数据行数增删模板数据行（无合并单元格，直接克隆），填文本。"""
    tbl = tpl_tb._tbl
    trs = tbl.findall(qn("w:tr"))
    proto_data = copy.deepcopy(trs[1])
    proto_total = copy.deepcopy(trs[-1])
    for tr in trs[1:]:
        tbl.remove(tr)
    for texts in data_rows:
        tr = copy.deepcopy(proto_data)
        fill_tr(tr, texts)
        tbl.append(tr)
    tr = copy.deepcopy(proto_total)
    if total_texts:
        fill_tr(tr, total_texts)
    tbl.append(tr)

    # 总价列过窄时数字会断行：任一总价超过 6 位字符则加宽（名称列让出 400 dxa）
    maxlen = max([len(t[3].strip()) for t in data_rows if len(t) > 3] +
                 ([len(total_texts[3].strip())] if total_texts and len(total_texts) > 3 else [0]))
    if maxlen > 6:
        grid = tbl.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol"))
        if len(cols) >= 4 and cols[3].get(qn("w:w")) == "1040":
            cols[0].set(qn("w:w"), str(int(cols[0].get(qn("w:w"))) - 400))
            cols[3].set(qn("w:w"), "1440")
            for tr in tbl.findall(qn("w:tr")):
                tcs = tr_tcs(tr)
                for ci in (0, 3):
                    if ci < len(tcs):
                        tcW = tcs[ci].find(qn("w:tcPr")).find(qn("w:tcW"))
                        if tcW is not None:
                            tcW.set(qn("w:w"), cols[ci].get(qn("w:w")))


# ---------------------------------------------------------------- 主流程

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("settlement", help="结算单 .docx")
    ap.add_argument("template", help="ZKHP 合同模板 .doc/.docx")
    ap.add_argument("-o", "--output-dir", default=".")
    ap.add_argument("--project", default=None, help="项目简称，如 绿洲15")
    ap.add_argument("--model", default=None, help="型号规格，如 YTS15-W-Z")
    ap.add_argument("--date", default=None, help="签订日期 YYYY-MM-DD，默认今天")
    ap.add_argument("--keep-number", action="store_true", help="保留模板原合同编号")
    args = ap.parse_args()

    settle_path = Path(args.settlement)
    date = datetime.strptime(args.date, "%Y-%m-%d") if args.date else datetime.now()
    date_iso = date.strftime("%Y-%m-%d")
    date_compact = date.strftime("%Y%m%d")

    project = args.project
    if not project:
        m = re.match(r"^\d{6}(.+?)(?:结算单|合同|[-_.]|$)", settle_path.stem)
        project = m.group(1) if m else ""
        print(f"[信息] 项目简称从文件名提取：{project or '(未提取到)'}")

    model = args.model
    if not model:
        m = re.search(r"(\d+)$", project)
        if m:
            model = f"YTS{m.group(1)}-W-Z"
            print(f"[信息] 型号由项目简称推导：{model}（如不对请用 --model 指定）")

    # 1. 读取结算单，算总价
    s_doc = docx.Document(str(settle_path))
    s_jiagong, s_dansuan = find_settlement_tables(s_doc)
    jg_total = table_total(s_jiagong, 9)
    ds_total = table_total(s_dansuan, 3)
    total_int = int((jg_total + ds_total).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    print(f"[信息] 加工表合计 {jg_total} + 单算表合计 {ds_total} = {jg_total + ds_total} → 合同总价 {total_int}")

    groups, jg_total_texts = classify_jiagong_rows(s_jiagong)
    _, ds_data, ds_total_texts = dansuan_data_rows(s_dansuan)

    # 2. 模板转 docx 并填充
    with tempfile.TemporaryDirectory() as td:
        tpl_docx = convert_doc_to_docx(Path(args.template), Path(td))
        t_doc = docx.Document(str(tpl_docx))

        old = {}
        for p in t_doc.paragraphs:
            if "合同编号" in p.text and not args.keep_number:
                old["编号"] = p.text.split("合同编号")[-1].strip("：: ")
                replace_in_paragraph(p, r"ZKHP\d{6,8}", f"ZKHP{date_compact}")
            if "签订时间" in p.text:
                old["签订时间"] = re.search(r"\d{4}-\d{2}-\d{2}", p.text)
                old["签订时间"] = old["签订时间"].group(0) if old["签订时间"] else ""
                replace_in_paragraph(p, r"\d{4}-\d{2}-\d{2}", date_iso)

        old.update(fill_product_table(t_doc.tables[0], model, total_int, project))
        rebuild_jiagong(t_doc.tables[2], groups, jg_total_texts)
        rebuild_dansuan(t_doc.tables[3], ds_data, ds_total_texts)

        out_name = f"{date.strftime('%y%m%d')}zkhp{project}合同{total_int}.docx"
        out_path = Path(args.output_dir) / out_name
        out_path.parent.mkdir(parents=True, exist_ok=True)
        t_doc.save(str(out_path))

    # 3. 校验：重读输出，报告旧值残留
    check = docx.Document(str(out_path))
    full = "\n".join(p.text for p in check.paragraphs)
    for tb in check.tables:
        for row in tb.rows:
            for c in row.cells:
                full += "\n" + c.text
    leftovers = [v for v in (old.get("price"), old.get("model"), old.get("remark"),
                             old.get("签订时间"), old.get("编号"))
                 if v and v in full]
    print(f"[完成] {out_path}")
    print(f"       签订时间 {date_iso}，合同编号 ZKHP{date_compact}，"
          f"型号 {model or '(保留模板)'}，备注 {project}，总价 {total_int}（{rmb_upper(total_int)}）")
    if leftovers:
        print(f"[警告] 以下旧值仍残留在文档中，请人工检查：{leftovers}")
    else:
        print("[校验] 无旧值残留")
    return str(out_path)


if __name__ == "__main__":
    main()
